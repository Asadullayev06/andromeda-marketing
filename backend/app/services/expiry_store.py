from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import os
from pathlib import Path
import re
from threading import Lock
from typing import BinaryIO

from ..config import settings

_lock = Lock()
_cached_mtime_ns: int | None = None
_cached_payload: dict | None = None
_cached_index_payload: dict | None = None
_cached_indexes: tuple[dict[str, list[dict]], dict[str, list[dict]]] | None = None


def normalize_product_name(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().casefold())


def snapshot_path() -> Path:
    return settings.data_dir / "company_stock_expiries.json"


def _payload() -> dict:
    global _cached_mtime_ns, _cached_payload
    path = snapshot_path()
    if not path.exists():
        path = Path(__file__).resolve().parent.parent / "data" / "company_stock_expiries.json"
    stat = path.stat()
    if _cached_payload is None or _cached_mtime_ns != stat.st_mtime_ns:
        _cached_payload = json.loads(path.read_text(encoding="utf-8"))
        _cached_mtime_ns = stat.st_mtime_ns
    return _cached_payload


def snapshot_info() -> dict:
    payload = _payload()
    items = payload.get("items", [])
    products = {(str(row.get("product_code") or ""), normalize_product_name(str(row.get("product_name") or ""))) for row in items}
    imported_at = date.fromisoformat(payload.get("imported_at"))
    age_days = max(0, (date.today() - imported_at).days)
    return {
        "source": payload.get("source", "Smartup stock detail export"),
        "imported_at": payload.get("imported_at"),
        "row_count": int(payload.get("raw_row_count") or len(items)),
        "aggregated_row_count": len(items),
        "product_count": len(products),
        "total_quantity": sum(float(row.get("quantity") or 0) for row in items),
        "age_days": age_days,
        "is_stale": age_days > 7,
    }


def indexes() -> tuple[dict[str, list[dict]], dict[str, list[dict]]]:
    global _cached_index_payload, _cached_indexes
    payload = _payload()
    if _cached_indexes is not None and _cached_index_payload is payload:
        return _cached_indexes
    by_code: dict[str, list[dict]] = {}
    by_name: dict[str, list[dict]] = {}
    for row in payload.get("items", []):
        code = str(row.get("product_code") or "").strip()
        name = normalize_product_name(str(row.get("product_name") or ""))
        if code:
            by_code.setdefault(code, []).append(row)
        if name:
            by_name.setdefault(name, []).append(row)
    _cached_index_payload = payload
    _cached_indexes = (by_code, by_name)
    return _cached_indexes


def rows_for_product(product) -> list[dict]:
    by_code, by_name = indexes()
    code = str(product.external_id or "").strip()
    if code and code in by_code:
        return by_code[code]
    return by_name.get(normalize_product_name(product.name), [])


def _date_value(value: object) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Invalid expiry date: {text}")


def import_workbook(content: bytes, filename: str) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel import dependency is not installed.") from exc
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Only .xlsx files are supported.")
    if len(content) > 20 * 1024 * 1024:
        raise ValueError("Workbook exceeds the 20 MB limit.")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header or len(header) < 6:
        raise ValueError("Expected six columns: warehouse, code, product, expiry, batch, available.")
    expected = (
        {"склад", "warehouse", "ombor"},
        {"код", "code", "kod"},
        {"название", "наименование", "product", "name", "mahsulot"},
        {"срок годности", "expiry", "expiration", "yaroqlilik muddati"},
        {"номер карточки", "batch", "card number", "партия", "partiya"},
        {"доступно", "available", "quantity", "miqdor"},
    )
    normalized_headers = [normalize_product_name(str(value or "")) for value in header[:6]]
    if any(value not in accepted for value, accepted in zip(normalized_headers, expected)):
        raise ValueError("Workbook columns must be: warehouse, code, product, expiry, batch/card number, available quantity.")

    aggregate: defaultdict[tuple[str, str, str | None, str | None], Decimal] = defaultdict(Decimal)
    raw_rows = 0
    errors: list[str] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = list(row) + [None] * max(0, 6 - len(row))
        code = str(values[1] or "").strip()
        name = str(values[2] or "").strip()
        if not name:
            continue
        try:
            expiry = _date_value(values[3])
            batch = str(values[4] or "").strip() or None
            quantity = Decimal(str(values[5] or 0))
            if quantity < 0:
                raise ValueError("quantity is negative")
        except (ValueError, InvalidOperation) as exc:
            if len(errors) < 20:
                errors.append(f"Row {row_number}: {exc}")
            continue
        aggregate[(code, name, expiry, batch)] += quantity
        raw_rows += 1
    workbook.close()
    if errors:
        raise ValueError("Workbook validation failed. " + "; ".join(errors))
    if not aggregate:
        raise ValueError("Workbook contains no usable stock rows.")

    items = [
        {
            "product_code": code,
            "product_name": name,
            "expiry_date": expiry,
            "batch_number": batch,
            "quantity": float(quantity),
        }
        for (code, name, expiry, batch), quantity in sorted(
            aggregate.items(), key=lambda item: (item[0][1].casefold(), item[0][2] or "9999", item[0][3] or "")
        )
    ]
    now = datetime.now(timezone.utc)
    payload = {
        "source": filename,
        "imported_at": now.date().isoformat(),
        "imported_at_utc": now.isoformat(),
        "raw_row_count": raw_rows,
        "items": items,
    }
    path = snapshot_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(".json.tmp")
    with _lock:
        temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(temp, path)
        global _cached_mtime_ns, _cached_payload, _cached_index_payload, _cached_indexes
        _cached_mtime_ns = None
        _cached_payload = None
        _cached_index_payload = None
        _cached_indexes = None
    return snapshot_info()
